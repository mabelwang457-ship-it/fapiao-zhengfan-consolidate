#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
农夫开票明细 · 正单/返单 按品类合计工具
========================================
功能：
  读取含「正单」「返单」两个 sheet 的同格式 xlsx，按【商品编码】聚合：
    - 正单逐行视为正值，返单逐行视为负值（返单数量/金额若为正会自动转负）
    - 同一商品编码的多行合并为一行：数量求和、金额(含税)求和
    - 产出「正单+返单_按品类合计」(主表) 与 「正单+返单_明细」(备查)
  净额 = 正单 + 返单（返单为负）；净额为负表示退货多于进货。

用法：
  python consolidate.py <输入.xlsx> [输出.xlsx]
  不传输出路径时，默认在输入文件同目录生成 <原名>_合并.xlsx

依赖：openpyxl（首次运行若缺失会自动尝试 pip install）
"""

import sys
import os


# ---------- 依赖自动安装 ----------
try:
    import openpyxl
except ImportError:
    import subprocess
    print("未检测到 openpyxl，正在尝试自动安装……")
    try:
        subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
        import openpyxl
    except Exception as e:
        print("自动安装失败，请手动执行：pip install openpyxl")
        print("错误：", e)
        sys.exit(1)


# ---------- 文本清洗 ----------
def clean(v):
    if isinstance(v, str):
        return v.replace('\xa0', '').replace('\u200b', '').strip()
    return v


def norm_code(v):
    if v is None:
        return None
    return str(v).replace('\xa0', '').replace(' ', '').strip()


def to_neg(v):
    """确保返单的数量/金额为负值（防御数据源正负号不一致）。"""
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
        return -v
    return v


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# ---------- sheet 定位 ----------
def find_sheet(wb, *keywords, exclude=None):
    for kw in keywords:
        for n in wb.sheetnames:
            if kw in n and (exclude is None or exclude not in n):
                return n
    return None


def resolve_columns(header):
    idx = {}
    for i, h in enumerate(header):
        hh = clean(h)
        if not hh:
            continue
        if '商品编码' in hh or hh == '编码':
            idx.setdefault('code', i)
        elif hh == '商品名称':
            idx.setdefault('name', i)
        elif hh == '规格型号':
            idx.setdefault('spec', i)
        elif hh == '计量单位':
            idx.setdefault('unit', i)
        elif '不含税价' in hh:
            idx.setdefault('price_no_tax', i)
        elif '含税价' in hh:
            idx.setdefault('price_tax', i)
        elif '数量' in hh:
            idx.setdefault('qty', i)
        elif '金额' in hh and '含税' in hh:
            idx.setdefault('amount', i)
    defaults = {'code': 0, 'name': 1, 'spec': 2, 'unit': 3,
                'qty': 4, 'price_no_tax': 5, 'price_tax': 6, 'amount': 7}
    for k, v in defaults.items():
        idx.setdefault(k, v)
    return idx


# ---------- 主流程 ----------
def consolidate(input_path, output_path=None):
    wb = openpyxl.load_workbook(input_path, data_only=True)

    zh_name = find_sheet(wb, '正单', exclude='+')
    fd_name = find_sheet(wb, '返单')
    if not zh_name:
        raise SystemExit("未找到「正单」sheet，请确认表格格式。")
    if not fd_name:
        raise SystemExit("未找到「返单」sheet，请确认表格格式。")

    zh_ws = wb[zh_name]
    zh_rows = [r for r in zh_ws.iter_rows(values_only=True)
               if any(c is not None for c in r)]
    zh_col = resolve_columns(zh_rows[0])
    zh_data = zh_rows[1:]

    fd_ws = wb[fd_name]
    fd_rows = [r for r in fd_ws.iter_rows(values_only=True)
               if any(c is not None for c in r)]
    fd_col = resolve_columns(fd_rows[0])
    fd_data = fd_rows[1:]

    summary = {}        # 按商品编码聚合
    detail_rows = []    # 明细（备查，正单+返单逐行）

    def ensure(code_raw, name, spec, unit):
        key = norm_code(code_raw)
        if key is None:
            return None
        s = summary.get(key)
        if s is None:
            s = {'code': clean(code_raw), 'name': clean(name),
                 'spec': clean(spec), 'unit': clean(unit),
                 'zq_qty': 0, 'fd_qty': 0, 'zq_amt': 0.0, 'fd_amt': 0.0}
            summary[key] = s
        return s

    # 正单
    for r in zh_data:
        code = clean(r[zh_col['code']])
        name = clean(r[zh_col['name']])
        spec = clean(r[zh_col['spec']])
        unit = clean(r[zh_col['unit']])
        qty = r[zh_col['qty']]
        pnt = r[zh_col['price_no_tax']]
        pt = r[zh_col['price_tax']]
        amt = r[zh_col['amount']]
        detail_rows.append([code, name, spec, unit, qty, pnt, pt, amt])
        s = ensure(code, name, spec, unit)
        if s is not None:
            if is_num(qty): s['zq_qty'] += qty
            if is_num(amt): s['zq_amt'] += amt

    # 返单（强制转负）
    for r in fd_data:
        code = clean(r[fd_col['code']])
        name = clean(r[fd_col['name']])
        spec = clean(r[fd_col['spec']])
        unit = clean(r[fd_col['unit']])
        qty = to_neg(r[fd_col['qty']])
        pnt = r[fd_col['price_no_tax']]
        pt = r[fd_col['price_tax']]
        amt = to_neg(r[fd_col['amount']])
        detail_rows.append([code, name, spec, unit, qty, pnt, pt, amt])
        s = ensure(code, name, spec, unit)
        if s is not None:
            if is_num(qty): s['fd_qty'] += qty
            if is_num(amt): s['fd_amt'] += amt

    # 构建聚合行（按编码排序，编码尽量按数值排）
    def sort_key(k):
        return (len(k), k)
    sum_header = ['商品编码', '商品名称', '规格型号', '计量单位',
                  '正单数量', '返单数量', '净数量',
                  '正单金额含税', '返单金额含税', '净金额含税', '平均含税单价(参考)']
    sum_rows = []
    for key in sorted(summary.keys(), key=sort_key):
        s = summary[key]
        net_qty = s['zq_qty'] + s['fd_qty']
        net_amt = round(s['zq_amt'] + s['fd_amt'], 2)
        avg = round(net_amt / net_qty, 2) if is_num(net_qty) and net_qty > 0 else None
        sum_rows.append([s['code'], s['name'], s['spec'], s['unit'],
                         s['zq_qty'], s['fd_qty'], net_qty,
                         round(s['zq_amt'], 2), round(s['fd_amt'], 2), net_amt, avg])

    detail_header = ['商品编码', '商品名称', '规格型号', '计量单位',
                     '入库数量', '不含税价', '含税价', '入库金额含税']

    # 写出（聚合表为主表，明细表为备查）
    out = openpyxl.Workbook()
    ws_sum = out.active
    ws_sum.title = '正单+返单_按品类合计'
    ws_sum.append(sum_header)
    for row in sum_rows:
        ws_sum.append(row)

    ws_det = out.create_sheet('正单+返单_明细')
    ws_det.append(detail_header)
    for row in detail_rows:
        ws_det.append(row)

    if output_path is None:
        base, ext = os.path.splitext(input_path)
        output_path = base + '_合并' + (ext or '.xlsx')
    out.save(output_path)

    print(f"输入：{input_path}")
    print(f"输出：{output_path}")
    print(f"正单行数：{len(zh_data)}，返单行数：{len(fd_data)}")
    print(f"按品类合计：{len(summary)} 个唯一商品编码")
    print(f"明细备查行数：{len(detail_rows)}")
    print("完成：主表「正单+返单_按品类合计」+ 备查「正单+返单_明细」已生成。")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("用法：python consolidate.py <输入.xlsx> [输出.xlsx]")
        sys.exit(1)
    consolidate(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
